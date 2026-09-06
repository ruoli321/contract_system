# ════════════════════════════════════════════════
# controllers — Excel 导出控制器（openpyxl）
# M17 · 合同台账导出
# ════════════════════════════════════════════════
import io
import logging
from datetime import datetime
from urllib.parse import quote

from odoo import http, fields, _
from odoo.http import request, Response
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)


def _get_openpyxl():
    """延迟导入 openpyxl，缺失时返回 None 让调用方降级"""
    try:
        import openpyxl  # noqa: F401
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
        return openpyxl, Font, Alignment, PatternFill, Border, Side, get_column_letter
    except ImportError:
        _logger.warning("openpyxl 未安装，Excel 导出功能不可用")
        return None, None, None, None, None, None, None


class ContractExportController(http.Controller):
    """合同台账 Excel 导出"""

    @http.route(
        "/contract/export/excel",
        type="http", auth="user", methods=["GET"],
    )
    def export_excel(self, **kwargs):
        """
        导出合同列表为 Excel。

        查询参数：
            type: 可选，合同类型过滤（purchase/sale/service/lease）
            state: 可选，状态过滤（draft/approving/seal/archived/void）
            date_from: 可选，签订日期起 YYYY-MM-DD
            date_to: 可选，签订日期止 YYYY-MM-DD
        """
        openpyxl, Font, Alignment, PatternFill, Border, Side, get_column_letter = _get_openpyxl()
        if openpyxl is None:
            raise UserError(_(
                "openpyxl 未安装，请在 Docker 中执行: "
                "pip install openpyxl"
            ))

        domain = self._build_domain(kwargs)
        contracts = request.env["contract.contract"].sudo().search(
            domain, order="date_signed desc, code"
        )

        if not contracts:
            raise UserError(_("没有符合条件的合同可导出"))

        wb = self._build_workbook(contracts, openpyxl, Font, Alignment, PatternFill, Border, Side, get_column_letter)

        # 输出到内存
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)

        filename = f"合同台账_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return request.make_response(
            stream.getvalue(),
            headers=[
                ("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                # HTTP 头仅支持 latin-1，中文文件名必须用 RFC 5987 的 filename* 形式
                ("Content-Disposition",
                 f"attachment; filename=\"contract_export.xlsx\"; filename*=UTF-8''{quote(filename)}"),
            ],
        )

    # ════════════════════════════════════════════════
    # 辅助
    # ════════════════════════════════════════════════
    def _build_domain(self, params):
        domain = []
        if params.get("type"):
            domain.append(("type", "=", params["type"]))
        if params.get("state"):
            domain.append(("state", "=", params["state"]))
        if params.get("date_from"):
            domain.append(("date_signed", ">=", params["date_from"]))
        if params.get("date_to"):
            domain.append(("date_signed", "<=", params["date_to"]))
        return domain

    def _build_workbook(self, contracts, openpyxl, Font, Alignment, PatternFill, Border, Side, get_column_letter):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "合同台账"

        # 样式
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        currency_fmt = '"¥"#,##0.00'
        center_align = Alignment(horizontal="center", vertical="center")

        # ── Sheet 1: 合同列表 ──
        headers = [
            "合同编号", "合同名称", "类型", "状态",
            "甲方", "乙方", "签订日期", "生效日期", "到期日期",
            "合同金额", "币种", "计划总额", "已完成", "剩余",
            "付款计划数", "逾期数", "AI 来源",
        ]
        state_map = dict(request.env["contract.contract"]._fields["state"]._description_selection(request.env))
        type_map = dict(request.env["contract.contract"]._fields["type"]._description_selection(request.env))

        # 写标题行 + 合并
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        title_cell = ws.cell(row=1, column=1, value=f"合同台账 ({len(contracts)} 条)")
        title_cell.font = Font(bold=True, size=14)
        title_cell.alignment = center_align

        # 写表头
        for col_idx, h in enumerate(headers, 1):
            c = ws.cell(row=3, column=col_idx, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = header_align
            c.border = thin_border

        # 写数据
        for row_idx, rec in enumerate(contracts, 4):
            ws.cell(row=row_idx, column=1, value=rec.code or "/").border = thin_border
            ws.cell(row=row_idx, column=2, value=rec.name or "").border = thin_border
            ws.cell(row=row_idx, column=3, value=type_map.get(rec.type, rec.type) or "").border = thin_border
            ws.cell(row=row_idx, column=4, value=state_map.get(rec.state, rec.state) or "").border = thin_border
            ws.cell(row=row_idx, column=5, value=rec.partner_a.name or "").border = thin_border
            ws.cell(row=row_idx, column=6, value=rec.partner_b.name or "").border = thin_border
            ws.cell(row=row_idx, column=7, value=str(rec.date_signed or "")).border = thin_border
            ws.cell(row=row_idx, column=8, value=str(rec.date_start or "")).border = thin_border
            ws.cell(row=row_idx, column=9, value=str(rec.date_end or "")).border = thin_border
            # 金额列带格式
            c_amt = ws.cell(row=row_idx, column=10, value=rec.amount or 0)
            c_amt.number_format = currency_fmt
            c_amt.border = thin_border
            ws.cell(row=row_idx, column=11, value=rec.currency_id.name or "").border = thin_border
            c_plan = ws.cell(row=row_idx, column=12, value=rec.planned_amount_total or 0)
            c_plan.number_format = currency_fmt
            c_plan.border = thin_border
            c_paid = ws.cell(row=row_idx, column=13, value=rec.paid_amount_total or 0)
            c_paid.number_format = currency_fmt
            c_paid.border = thin_border
            c_rem = ws.cell(row=row_idx, column=14, value=rec.remaining_amount_total or 0)
            c_rem.number_format = currency_fmt
            c_rem.border = thin_border
            ws.cell(row=row_idx, column=15, value=rec.payment_plan_count or 0).border = thin_border
            ws.cell(row=row_idx, column=16, value=rec.overdue_count or 0).border = thin_border
            ws.cell(row=row_idx, column=17, value="是" if rec.is_ai_generated else "否").border = thin_border

        # 自动列宽（简单估算）
        col_widths = [14, 40, 12, 10, 22, 22, 12, 12, 12, 14, 8, 14, 14, 14, 10, 8, 10]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        # ── Sheet 2: 汇总 ──
        self._write_summary_sheet(wb, contracts, Font, Alignment, PatternFill, Border, Side, get_column_letter, currency_fmt)

        return wb

    def _write_summary_sheet(self, wb, contracts, Font, Alignment, PatternFill, Border, Side, get_column_letter, currency_fmt):
        from collections import defaultdict
        ws = wb.create_sheet("汇总分析")

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

        type_map = dict(contracts[0]._fields["type"]._description_selection(request.env)) if contracts else {}
        state_map = dict(contracts[0]._fields["state"]._description_selection(request.env)) if contracts else {}

        # 按类型汇总
        ws.cell(row=1, column=1, value="一、按合同类型汇总").font = Font(bold=True, size=13)
        ws.cell(row=1, column=1).border = thin_border

        type_group = defaultdict(lambda: {"count": 0, "amount": 0, "paid": 0})
        for c in contracts:
            key = type_map.get(c.type, c.type or "未知")
            type_group[key]["count"] += 1
            type_group[key]["amount"] += c.amount or 0
            type_group[key]["paid"] += c.paid_amount_total or 0

        headers = ["类型", "合同数", "合同总额", "已完成金额", "已完成占比"]
        for col_idx, h in enumerate(headers, 1):
            c = ws.cell(row=2, column=col_idx, value=h)
            c.font = header_font; c.fill = header_fill; c.alignment = header_align; c.border = thin_border

        row = 3
        for key, v in sorted(type_group.items()):
            ws.cell(row=row, column=1, value=key).border = thin_border
            ws.cell(row=row, column=2, value=v["count"]).border = thin_border
            c_amt = ws.cell(row=row, column=3, value=v["amount"])
            c_amt.number_format = currency_fmt; c_amt.border = thin_border
            c_paid = ws.cell(row=row, column=4, value=v["paid"])
            c_paid.number_format = currency_fmt; c_paid.border = thin_border
            pct = (v["paid"] / v["amount"] * 100) if v["amount"] else 0
            c_pct = ws.cell(row=row, column=5, value=f"{pct:.1f}%")
            c_pct.border = thin_border; c_pct.alignment = header_align
            row += 1

        # 合计行
        ws.cell(row=row, column=1, value="合计").font = Font(bold=True)
        ws.cell(row=row, column=1).border = thin_border
        total_count = sum(v["count"] for v in type_group.values())
        total_amt = sum(v["amount"] for v in type_group.values())
        total_paid = sum(v["paid"] for v in type_group.values())
        ws.cell(row=row, column=2, value=total_count).border = thin_border
        ws.cell(row=row, column=2).font = Font(bold=True)
        c = ws.cell(row=row, column=3, value=total_amt); c.number_format = currency_fmt; c.font = Font(bold=True); c.border = thin_border
        c = ws.cell(row=row, column=4, value=total_paid); c.number_format = currency_fmt; c.font = Font(bold=True); c.border = thin_border
        pct = (total_paid / total_amt * 100) if total_amt else 0
        c = ws.cell(row=row, column=5, value=f"{pct:.1f}%"); c.font = Font(bold=True); c.border = thin_border; c.alignment = header_align

        ws.column_dimensions["A"].width = 16
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 14
